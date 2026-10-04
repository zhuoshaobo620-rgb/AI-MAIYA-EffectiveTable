# 作用：DATASET_SCOPE_ZERO_EXACT 契约——source 有行但正式深圳 dataset scope 内精确为 0（非 source-native zero）。
# 何时改：与 daily_report.process_exported_file / VALID_PLATFORMS 口径同步时；禁止为单日硬编码 PASS。

from __future__ import annotations

import hashlib
import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

from crm_export_artifact_gate import (
    CRM_EXPORT_REQUIRED_COLUMNS,
    EXPORT_PROCESS_ZERO_YIELD,
    validate_crm_export_structure,
)

DATASET_SCOPE_ZERO_CONTRACT_V2 = "DATASET_SCOPE_ZERO_CONTRACT_V2"
DATASET_SCOPE_ZERO_EVIDENCE_FILENAME = "crm_dataset_scope_zero_evidence.json"
SCOPE_ARTIFACT_KIND = "DATASET_SCOPE_ZERO_EVIDENCE_ARTIFACT"

SCOPE_SOURCE_OF_TRUTH = "daily_report.VALID_PLATFORMS+process_exported_file"


def is_dataset_scope_zero_evidence_path(path: str | Path) -> bool:
    return Path(path).name == DATASET_SCOPE_ZERO_EVIDENCE_FILENAME


def _parse_date_value(值: object) -> date | None:
    if 值 is None:
        return None
    if isinstance(值, datetime):
        return 值.date()
    if isinstance(值, date):
        return 值
    文本 = str(值).strip()
    if len(文本) < 10:
        return None
    try:
        return datetime.strptime(文本[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def _row_id_hash(row_index: int, platform: str) -> str:
    载荷 = f"row:{row_index}|platform:{platform.strip()}"
    return hashlib.sha256(载荷.encode("utf-8")).hexdigest()[:16]


def audit_crm_export_scope_rows(
    export_path: str | Path,
    business_date: str,
    *,
    valid_platforms: tuple[str, ...] | None = None,
) -> dict[str, Any]:
    """
    行级 scope 审计（禁止输出 PII；仅 row_id_hash + platform + classification）。
    分类规则与 process_exported_file 一致。
    """
    from daily_report import CRM_COLUMNS, VALID_PLATFORMS, parse_datetime_value

    白名单 = set(valid_platforms or VALID_PLATFORMS)
    目标日 = date.fromisoformat(str(business_date)[:10])
    路径 = Path(export_path)
    行审计: list[dict[str, Any]] = []
    parser_error_rows = 0
    unknown_rows = 0
    in_scope = 0
    out_of_scope = 0

    if not 路径.is_file():
        return {
            "source_row_count": 0,
            "in_scope_row_count": 0,
            "out_of_scope_row_count": 0,
            "unknown_row_count": 1,
            "parser_error_rows": 1,
            "unclassified_rows": 1,
            "row_classifications": [],
            "period_exact": False,
        }

    wb = openpyxl.load_workbook(路径, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    try:
        col_idx = {name: headers.index(name) + 1 for name in CRM_COLUMNS}
    except ValueError:
        wb.close()
        return {
            "source_row_count": 0,
            "in_scope_row_count": 0,
            "out_of_scope_row_count": 0,
            "unknown_row_count": 0,
            "parser_error_rows": 1,
            "unclassified_rows": 1,
            "row_classifications": [],
            "period_exact": True,
        }

    for row_index, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
        if not any(cell.value not in (None, "") for cell in row):
            continue
        try:
            values = {name: row[col_idx[name] - 1].value for name in CRM_COLUMNS}
        except Exception:
            parser_error_rows += 1
            unknown_rows += 1
            continue
        平台 = str(values.get("来源平台") or "").strip().replace("　", " ").strip()
        行号哈希 = _row_id_hash(row_index, 平台)
        if "香港" in 平台:
            out_of_scope += 1
            行审计.append(
                {
                    "row_id_hash": 行号哈希,
                    "source_platform": 平台,
                    "normalized_platform": 平台,
                    "scope_membership": "OUT_OF_SCOPE",
                    "classification": "hong_kong_excluded",
                }
            )
            continue
        合法平台 = 平台 in 白名单 or "三方" in 平台
        if not 合法平台:
            out_of_scope += 1
            行审计.append(
                {
                    "row_id_hash": 行号哈希,
                    "source_platform": 平台,
                    "normalized_platform": 平台,
                    "scope_membership": "OUT_OF_SCOPE",
                    "classification": "platform_not_in_valid_scope",
                }
            )
            continue
        if "三方数据" in 平台 and "3" in 平台 and "惠州" not in 平台:
            标签 = str(values.get("客户标签") or "")
            if 标签 == "重复":
                out_of_scope += 1
                行审计.append(
                    {
                        "row_id_hash": 行号哈希,
                        "source_platform": 平台,
                        "normalized_platform": 平台,
                        "scope_membership": "OUT_OF_SCOPE",
                        "classification": "sanfang_duplicate_label",
                    }
                )
                continue
            表单 = str(values.get("表单名称") or "")
            线下 = "-线下" in 表单
            建档 = parse_datetime_value(values.get("建档时间"))
            if not 线下 and (建档 is None or 建档.date() != 目标日):
                out_of_scope += 1
                行审计.append(
                    {
                        "row_id_hash": 行号哈希,
                        "source_platform": 平台,
                        "normalized_platform": 平台,
                        "scope_membership": "OUT_OF_SCOPE",
                        "classification": "sanfang_archive_date_not_target_day",
                    }
                )
                continue
            in_scope += 1
            行审计.append(
                {
                    "row_id_hash": 行号哈希,
                    "source_platform": 平台,
                    "normalized_platform": 平台,
                    "scope_membership": "IN_SCOPE",
                    "classification": "retained_sanfang",
                }
            )
            continue
        建档 = parse_datetime_value(values.get("建档时间"))
        if 建档 is None or 建档.date() != 目标日:
            out_of_scope += 1
            行审计.append(
                {
                    "row_id_hash": 行号哈希,
                    "source_platform": 平台,
                    "normalized_platform": 平台,
                    "scope_membership": "OUT_OF_SCOPE",
                    "classification": "archive_date_not_target_day_or_duplicate",
                }
            )
            continue
        in_scope += 1
        行审计.append(
            {
                "row_id_hash": 行号哈希,
                "source_platform": 平台,
                "normalized_platform": 平台,
                "scope_membership": "IN_SCOPE",
                "classification": "retained_default_platform",
            }
        )

    wb.close()
    源行 = in_scope + out_of_scope + unknown_rows
    return {
        "source_row_count": 源行,
        "in_scope_row_count": in_scope,
        "out_of_scope_row_count": out_of_scope,
        "unknown_row_count": unknown_rows,
        "parser_error_rows": parser_error_rows,
        "unclassified_rows": unknown_rows + parser_error_rows,
        "row_classifications": 行审计,
        "period_exact": True,
        "scope_source": SCOPE_SOURCE_OF_TRUTH,
        "scope_contract_version": DATASET_SCOPE_ZERO_CONTRACT_V2,
    }


def evaluate_dataset_scope_zero_exact(
    export_path: str | Path,
    business_date: str,
    *,
    structure_gate: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """fail-closed：仅当 source>0、结构有效、全部行可分类且 in_scope=0 时 PASS。"""
    路径 = Path(export_path)
    结构 = structure_gate or validate_crm_export_structure(路径)
    审计 = audit_crm_export_scope_rows(路径, business_date)
    失败: list[str] = []
    if 结构.get("EXPORT_ARTIFACT_STRUCTURE_VALID") != "PASS":
        失败.append("export_structure_invalid")
    源行 = int(审计.get("source_row_count") or 0)
    if 源行 <= 0:
        失败.append("source_row_count_not_positive")
    if int(审计.get("in_scope_row_count") or 0) != 0:
        失败.append("in_scope_rows_present")
    if int(审计.get("unknown_row_count") or 0) != 0:
        失败.append("unknown_rows")
    if int(审计.get("parser_error_rows") or 0) != 0:
        失败.append("parser_error_rows")
    if int(审计.get("out_of_scope_row_count") or 0) != 源行:
        失败.append("out_of_scope_rows_mismatch")
    指纹 = ""
    if 路径.is_file():
        指纹 = hashlib.sha256(路径.read_bytes()).hexdigest()
    通过 = not 失败
    return {
        "DATASET_SCOPE_ZERO_EXACT": "PASS" if 通过 else "BLOCKED",
        "contract_version": DATASET_SCOPE_ZERO_CONTRACT_V2,
        "scope_source": SCOPE_SOURCE_OF_TRUTH,
        "period_start": str(business_date)[:10],
        "period_end": str(business_date)[:10],
        "source_artifact_sha256": 指纹,
        "export_structure_valid": 结构.get("EXPORT_ARTIFACT_STRUCTURE_VALID") == "PASS",
        "failures": 失败 or None,
        "audit": 审计,
        "zero_semantics": "dataset_scope_zero_exact",
        "export_process_zero_yield_would_apply": bool(
            结构.get("EXPORT_ARTIFACT_STRUCTURE_VALID") == "PASS"
            and 源行 > 0
            and int(审计.get("in_scope_row_count") or 0) == 0
            and not 通过
        ),
        "code": None if 通过 else EXPORT_PROCESS_ZERO_YIELD,
    }


def write_dataset_scope_zero_evidence(
    download_dir: str | Path,
    *,
    business_date: str,
    export_path: str | Path,
    evaluation: dict[str, Any],
) -> Path:
    目录 = Path(download_dir)
    目录.mkdir(parents=True, exist_ok=True)
    路径 = 目录 / DATASET_SCOPE_ZERO_EVIDENCE_FILENAME
    审计 = evaluation.get("audit") or {}
    载荷 = {
        "artifact_kind": SCOPE_ARTIFACT_KIND,
        "contract_version": DATASET_SCOPE_ZERO_CONTRACT_V2,
        "dataset_scope_zero_exact": evaluation.get("DATASET_SCOPE_ZERO_EXACT") == "PASS",
        "zero_semantics": "dataset_scope_zero_exact",
        "business_date": str(business_date)[:10],
        "period_start": str(business_date)[:10],
        "period_end": str(business_date)[:10],
        "source_row_count": 审计.get("source_row_count"),
        "in_scope_row_count": 审计.get("in_scope_row_count"),
        "out_of_scope_row_count": 审计.get("out_of_scope_row_count"),
        "unknown_row_count": 审计.get("unknown_row_count"),
        "parser_error_rows": 审计.get("parser_error_rows"),
        "scope_source": 审计.get("scope_source"),
        "scope_contract_version": 审计.get("scope_contract_version"),
        "source_artifact_sha256": evaluation.get("source_artifact_sha256"),
        "source_artifact_basename": Path(export_path).name,
        "row_classifications": 审计.get("row_classifications") or [],
        "technical_failure_flag": False,
    }
    路径.write_text(json.dumps(载荷, ensure_ascii=False, indent=2), encoding="utf-8")
    return 路径
