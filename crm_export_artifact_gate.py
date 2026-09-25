# 作用：CRM 信息流线索池导出 xlsx 结构校验与下载目录候选选择（禁止 read_only 误判行数）。
# 何时改：CRM 导出列契约或 daily_report 处理前门禁规则变化时。

from __future__ import annotations

import os
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl

# 与 daily_report.CRM_COLUMNS 保持一致
CRM_EXPORT_REQUIRED_COLUMNS = (
    "客户姓名",
    "电话",
    "记录时间",
    "来源平台",
    "表单名称",
    "建档时间",
    "客户标签",
    "网电咨询师",
)

WRONG_OR_MALFORMED_EXPORT_ARTIFACT = "WRONG_OR_MALFORMED_EXPORT_ARTIFACT"
EXPORT_PROCESS_ZERO_YIELD = "EXPORT_PROCESS_ZERO_YIELD"

# 单列「记录时间」表头且无数据行 → 典型错误导出变体
_MALFORMED_SINGLE_HEADER = ("记录时间",)


@dataclass(frozen=True)
class ExportArtifactInspect:
    sheet_names: tuple[str, ...]
    active_sheet: str
    max_row: int
    max_column: int
    headers: tuple[str, ...]
    data_row_count: int
    required_columns_present: bool
    missing_columns: tuple[str, ...]


def _load_workbook_full(path: Path) -> openpyxl.Workbook:
    return openpyxl.load_workbook(path, data_only=True)


def inspect_crm_export_artifact(path: str | Path) -> ExportArtifactInspect:
    """完整加载工作簿统计行数（禁止 read_only，避免 CRM 导出 max_row=1 误判）。。"""
    文件 = Path(path)
    wb = _load_workbook_full(文件)
    ws = wb.active
    headers = tuple(str(cell.value or "") for cell in next(ws.iter_rows(min_row=1, max_row=1)))
    missing = tuple(c for c in CRM_EXPORT_REQUIRED_COLUMNS if c not in headers)
    data_rows = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(v not in (None, "") for v in row):
            data_rows += 1
    result = ExportArtifactInspect(
        sheet_names=tuple(wb.sheetnames),
        active_sheet=str(ws.title),
        max_row=int(ws.max_row or 0),
        max_column=int(ws.max_column or 0),
        headers=headers,
        data_row_count=data_rows,
        required_columns_present=len(missing) == 0,
        missing_columns=missing,
    )
    wb.close()
    return result


def validate_crm_export_structure(path: str | Path) -> dict[str, Any]:
    """
    EXPORT_ARTIFACT_STRUCTURE_VALID 门禁。
    FAIL 时 code=WRONG_OR_MALFORMED_EXPORT_ARTIFACT。
    """
    文件 = Path(path)
    if not 文件.is_file():
        return {
            "EXPORT_ARTIFACT_STRUCTURE_VALID": "FAIL",
            "code": WRONG_OR_MALFORMED_EXPORT_ARTIFACT,
            "reason": "export_file_missing",
            "path": str(文件),
        }
    try:
        检查 = inspect_crm_export_artifact(文件)
    except Exception as 错误:
        return {
            "EXPORT_ARTIFACT_STRUCTURE_VALID": "FAIL",
            "code": WRONG_OR_MALFORMED_EXPORT_ARTIFACT,
            "reason": f"open_failed:{错误}",
            "path": str(文件),
        }

    if not 检查.required_columns_present:
        return {
            "EXPORT_ARTIFACT_STRUCTURE_VALID": "FAIL",
            "code": WRONG_OR_MALFORMED_EXPORT_ARTIFACT,
            "reason": "missing_required_columns",
            "missing_columns": list(检查.missing_columns),
            "headers": list(检查.headers),
            "path": str(文件),
        }

    if 检查.max_column < 8:
        return {
            "EXPORT_ARTIFACT_STRUCTURE_VALID": "FAIL",
            "code": WRONG_OR_MALFORMED_EXPORT_ARTIFACT,
            "reason": "too_few_columns",
            "max_column": 检查.max_column,
            "path": str(文件),
        }

    if 检查.data_row_count < 1:
        if list(检查.headers) == list(_MALFORMED_SINGLE_HEADER) or (
            len(检查.headers) == 1 and 检查.headers[0] == "记录时间"
        ):
            reason = "malformed_single_column_record_time_header_only"
        else:
            reason = "no_data_rows"
        return {
            "EXPORT_ARTIFACT_STRUCTURE_VALID": "FAIL",
            "code": WRONG_OR_MALFORMED_EXPORT_ARTIFACT,
            "reason": reason,
            "max_row": 检查.max_row,
            "headers": list(检查.headers),
            "path": str(文件),
        }

    return {
        "EXPORT_ARTIFACT_STRUCTURE_VALID": "PASS",
        "code": None,
        "data_row_count": 检查.data_row_count,
        "max_row": 检查.max_row,
        "max_column": 检查.max_column,
        "sheet_names": list(检查.sheet_names),
        "headers_sample": list(检查.headers[:12]),
        "path": str(文件),
    }


def select_fresh_crm_export_xlsx(
    download_dir: str | Path,
    *,
    not_before_mtime: float | None = None,
    preferred_basenames: tuple[str, ...] = ("信息流查询.xlsx",),
) -> Path | None:
    """
    在下载目录中选择本次导出产物：优先 preferred 文件名，否则取 mtime 最新且 >= not_before_mtime 的 xlsx。
    排除含「建档」的 xlsx（建档池另选）。
    """
    目录 = Path(download_dir)
    if not 目录.is_dir():
        return None
    候选: list[Path] = []
    for 名称 in preferred_basenames:
        路径 = 目录 / 名称
        if 路径.is_file():
            候选.append(路径)
    if not 候选:
        for 路径 in 目录.glob("*.xlsx"):
            if "建档" in 路径.name:
                continue
            候选.append(路径)
    if not 候选:
        return None

    def _可接受(路径: Path) -> bool:
        if not_before_mtime is None:
            return True
        try:
            return os.path.getmtime(路径) >= not_before_mtime - 0.05
        except OSError:
            return False

    可选 = [p for p in 候选 if _可接受(p)]
    if not 可选:
        return None
    return max(可选, key=lambda p: os.path.getmtime(p))


def assert_export_process_yield(
    *,
    structure_gate: dict[str, Any],
    processed_row_count: int,
    jianDang_row_count: int = 0,
) -> dict[str, Any] | None:
    """
    导出结构 PASS 但业务过滤后 0 行 → EXPORT_PROCESS_ZERO_YIELD（避免与 source-empty 混淆）。
    """
    if structure_gate.get("EXPORT_ARTIFACT_STRUCTURE_VALID") != "PASS":
        return None
    raw = int(structure_gate.get("data_row_count") or 0)
    if raw > 0 and processed_row_count + jianDang_row_count == 0:
        return {
            "EXPORT_ARTIFACT_STRUCTURE_VALID": "PASS",
            "EXPORT_PROCESS_YIELD": "FAIL",
            "code": EXPORT_PROCESS_ZERO_YIELD,
            "reason": "crm_export_has_rows_but_none_pass_platform_filters",
            "raw_data_row_count": raw,
        }
    return None


SOURCE_NATIVE_ZERO_EVIDENCE_FILENAME = "crm_source_native_zero_evidence.json"
SOURCE_NATIVE_ZERO_EVIDENCE_CONTRACT = "SOURCE_NATIVE_ZERO_EVIDENCE_V1"
ZERO_EVIDENCE_ARTIFACT_KIND = "ZERO_EVIDENCE_ARTIFACT"


def is_source_native_zero_evidence_path(path: str | Path) -> bool:
    return Path(path).name == SOURCE_NATIVE_ZERO_EVIDENCE_FILENAME


def write_source_native_zero_evidence(
    download_dir: str | Path,
    *,
    business_date: str,
    ui_start_readback: str,
    ui_end_readback: str,
    query_timestamp: str,
    ui_state: str,
    record_count_evidence: int,
    screenshot_fingerprint_sha256: str | None = None,
    refresh_evidence: dict[str, Any] | None = None,
) -> Path:
    """写入确定性 source-native zero 证据 JSON（非伪装为非空 CRM xlsx）。"""
    目录 = Path(download_dir)
    目录.mkdir(parents=True, exist_ok=True)
    路径 = 目录 / SOURCE_NATIVE_ZERO_EVIDENCE_FILENAME
    载荷 = {
        "artifact_kind": ZERO_EVIDENCE_ARTIFACT_KIND,
        "contract_version": SOURCE_NATIVE_ZERO_EVIDENCE_CONTRACT,
        "source_native_zero": True,
        "raw_row_count": 0,
        "business_date": str(business_date)[:10],
        "requested_period_start": str(business_date)[:10],
        "requested_period_end": str(business_date)[:10],
        "ui_start_readback": ui_start_readback,
        "ui_end_readback": ui_end_readback,
        "query_timestamp": query_timestamp,
        "ui_state": ui_state,
        "record_count_evidence": int(record_count_evidence),
        "screenshot_fingerprint_sha256": screenshot_fingerprint_sha256,
        "refresh_evidence": refresh_evidence or {},
        "technical_failure_flag": False,
    }
    import json

    路径.write_text(json.dumps(载荷, ensure_ascii=False, indent=2), encoding="utf-8")
    return 路径


def load_source_native_zero_evidence(path: str | Path) -> dict[str, Any]:
    import json

    return json.loads(Path(path).read_text(encoding="utf-8"))


def validate_source_native_zero_evidence(payload: dict[str, Any]) -> dict[str, Any]:
    """评估证据包是否满足 source-native explicit zero（fail-closed）。"""
    失败: list[str] = []
    if payload.get("artifact_kind") != ZERO_EVIDENCE_ARTIFACT_KIND:
        失败.append("artifact_kind_mismatch")
    if payload.get("contract_version") != SOURCE_NATIVE_ZERO_EVIDENCE_CONTRACT:
        失败.append("contract_version_mismatch")
    if payload.get("source_native_zero") is not True:
        失败.append("source_native_zero_flag_missing")
    if payload.get("raw_row_count") is None or int(payload.get("raw_row_count")) != 0:
        失败.append("raw_row_count_not_zero")
    日 = str(payload.get("business_date") or "")[:10]
    开始 = str(payload.get("ui_start_readback") or "")[:10]
    结束 = str(payload.get("ui_end_readback") or "")[:10]
    if not 日 or 开始 != 日 or 结束 != 日:
        失败.append("period_not_exact")
    if payload.get("technical_failure_flag"):
        失败.append("technical_failure_flag_set")
    if payload.get("record_count_evidence") is None or int(payload.get("record_count_evidence")) != 0:
        失败.append("record_count_evidence_not_zero")
    ui_state = str(payload.get("ui_state") or "")
    if ui_state not in {"UI_EXPLICIT_ZERO", "SOURCE_NATIVE_ZERO_CONFIRMED"}:
        失败.append("ui_state_not_explicit_zero")
    刷新 = payload.get("refresh_evidence") or {}
    if 刷新.get("query_refresh_confirmed") is not True:
        失败.append("query_refresh_not_confirmed")
    if 刷新.get("stale_ui_risk") is True:
        失败.append("stale_ui_risk")
    有效 = not 失败
    return {
        "SOURCE_NATIVE_ZERO_EVIDENCE_VALID": "PASS" if 有效 else "BLOCKED",
        "evidence_fingerprint_sha256": payload.get("evidence_fingerprint_sha256"),
        "failures": 失败 or None,
        "contract_version": SOURCE_NATIVE_ZERO_EVIDENCE_CONTRACT,
    }
